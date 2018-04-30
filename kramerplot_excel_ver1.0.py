import pandas as pd 
import numpy as np
import openpyxl as px
from openpyxl.chart import (
    ScatterChart,
    Reference,
    Series,
)

#Excel ファイルを作成
wb = px.Workbook()
ws = wb.active

#CSVからデータを読み込む
data = pd.read_csv("data.csv", header = None,skiprows= 1)


num = len(data.columns)#データの列の数を取得
B1 = np.arange(8,27)#任意の磁場を設定
out = B1

#kramer plotの定数を設定
p = 0.5
q = 2

for i in range(1, num) :
    df = data.loc[:, [0,i]] #データを一つずつデータフレームに変換
    
    df = df.dropna()
    B = df[0]
    Jc = df[i]

    B = B.as_matrix()
    Jc = Jc.as_matrix()
    BJc = (B**0.25)*(Jc**0.5)
    #最小二乗解を計算
    A = np.array([B, np.ones(len(B))])
    A = A.T
    a,b = np.linalg.lstsq(A,BJc)[0]
    Bc2 = -b/a
    b1 = B1/Bc2
    b2 = B/Bc2
    C = Jc / ((b2**(p-1))*((1-b2)**q))
    C = np.average(C)
    Jcfit = C * (b1**(p-1))*((1-b1)**q)
    out = np.c_[out, Jcfit]

#np.savetxt('out.csv', out, delimiter=',') #csvファイルに保存



row_num = out.shape[0] #出力ファイルの行数を取得
col_num = out.shape[1] #出力ファイルの列数を取得

#データのタイトルを取得
data2 = pd.read_csv("data.csv", header = None)
title_data = data2.loc[0, :]
title_data = title_data.as_matrix()
title_list = title_data.tolist()

ws.append(title_list)

#ws.append(title_data)
out_list = out.tolist()

#Excelファイルへデータ書き込み
for j in out_list:
    ws.append(j)
#for j in range(1,row_num+1):
#    for k in range(1, col_num+1):
#        ws.cell(row = j, column = k, value = out[j-1, k-1])


##グラフの描画
chart = ScatterChart() #グラフの種類を選択
chart.title = "Jc-B performance"#グラフタイトル
chart.style = 13
chart.x_axis.title = 'Magnetic Field (T)'#x軸ラベル
chart.y_axis.title = 'Matrix Jc (A/mm^2)'#y軸ラベル




#x軸値の設定
xvalues = Reference(ws, min_col = 1, min_row = 1, max_row = row_num)
for w in range(2, col_num + 1):
    values = Reference(ws, min_col = w, min_row = 1, max_row = row_num)
    series = Series(values, xvalues, title_from_data = True)
    chart.series.append(series)



ws.add_chart(chart, "H1")
wb.save('kramerplot_out.xlsx')




